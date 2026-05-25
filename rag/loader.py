from __future__ import annotations

import hashlib
import os
import tempfile
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class TextBlock:
    kind: str
    text: str
    level: int | None = None


@dataclass(frozen=True)
class ParsedDocument:
    file_name: str
    sha256: str
    blocks: list[TextBlock]


SUPPORTED_EXTENSIONS = (".docx", ".md", ".txt", ".pdf")


def _normalize(text: str) -> str:
    return " ".join(text.split()).strip()


def _parse_docx(content: bytes) -> list[TextBlock]:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name

    try:
        doc = Document(tmp_path)
        blocks: list[TextBlock] = []

        for para in doc.paragraphs:
            text = _normalize(para.text or "")
            if not text:
                continue
            style_name = (para.style.name or "").lower() if para.style else ""
            if style_name.startswith("heading"):
                level = None
                parts = style_name.split()
                if len(parts) >= 2 and parts[1].isdigit():
                    level = int(parts[1])
                blocks.append(TextBlock(kind="heading", text=text, level=level))
            else:
                blocks.append(TextBlock(kind="paragraph", text=text))

        return blocks
    finally:
        os.unlink(tmp_path)


def _parse_markdown(content: bytes) -> list[TextBlock]:
    blocks: list[TextBlock] = []
    for line in content.decode("utf-8", errors="ignore").splitlines():
        text = line.strip()
        if not text:
            continue
        if text.startswith("#"):
            heading_text = _normalize(text.lstrip("#").strip())
            if not heading_text:
                continue
            level = len(text) - len(text.lstrip("#"))
            blocks.append(TextBlock(kind="heading", text=heading_text, level=level))
        else:
            blocks.append(TextBlock(kind="paragraph", text=_normalize(text)))
    return blocks


def _parse_text(content: bytes) -> list[TextBlock]:
    blocks: list[TextBlock] = []
    for line in content.decode("utf-8", errors="ignore").splitlines():
        text = _normalize(line)
        if text:
            blocks.append(TextBlock(kind="paragraph", text=text))
    return blocks


def _parse_pdf(content: bytes) -> list[TextBlock]:
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name

    try:
        blocks: list[TextBlock] = []
        reader = PdfReader(tmp_path)
        for page_num, page in enumerate(reader.pages, start=1):
            text = _normalize(page.extract_text() or "")
            if not text:
                continue
            blocks.append(TextBlock(kind="heading", text=f"Page {page_num}", level=2))
            blocks.append(TextBlock(kind="paragraph", text=text))
        return blocks
    finally:
        os.unlink(tmp_path)


def parse_uploaded_file(file_name: str, content: bytes) -> ParsedDocument:
    ext = os.path.splitext(file_name)[1].lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported: {ext}. Allowed: {', '.join(SUPPORTED_EXTENSIONS)}")

    parsers = {
        ".docx": lambda: _parse_docx(content),
        ".md": lambda: _parse_markdown(content),
        ".txt": lambda: _parse_text(content),
        ".pdf": lambda: _parse_pdf(content),
    }

    blocks = parsers[ext]()
    if not blocks:
        raise ValueError(f"Document is empty: {file_name}")

    return ParsedDocument(
        file_name=file_name,
        sha256=hashlib.sha256(content).hexdigest(),
        blocks=blocks,
    )
